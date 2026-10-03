"""Build editable Word manuscripts from reviewed Markdown, citations and figures.

Display mathematics are genuine OMML, never raster images. Unrecognized display
equations fail closed. The v1.9 scientific sources must be frozen before use.
Rendering requires an explicit runtime decision; no desktop LibreOffice fallback.
"""
from __future__ import annotations
import argparse
import copy
import hashlib
import json
import re
import subprocess
import sys
import os
from pathlib import Path
from zipfile import ZipFile
from datetime import datetime, timezone
from lxml import etree
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

D19 = Path('D:/论文/SCI投稿汇总/Onion_Deep_Revision_20261003_v19')
SKILL = Path('C:/Users/lixin/.codex/plugins/cache/openai-primary-runtime/documents/26.909.12148/skills/documents')
RUNTIME = Path('C:/Users/lixin/.cache/codex-runtimes/codex-primary-runtime')
MNS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
WNS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'m': MNS, 'w': WNS}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def el(name, value=None):
    n = OxmlElement(name)
    if value is not None:
        n.set(qn(name.split(':')[0] + ':val'), str(value))
    return n


def mathrun(text, plain=False):
    r = el('m:r')
    if plain:
        props = el('m:rPr'); props.append(el('m:sty', 'p')); r.append(props)
    wp = el('w:rPr'); font = el('w:rFonts')
    for name in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        font.set(qn('w:' + name), 'Cambria Math')
    wp.append(font); wp.append(el('w:sz', '22')); r.append(wp)
    t = el('m:t'); t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    r.append(t)
    return r


def box(name, children):
    x = el('m:' + name)
    for c in children:
        x.append(c)
    return x


def sub(base, lower, upper=None):
    if upper is None:
        return box('sSub', [box('e', base), box('sub', lower)])
    return box('sSubSup', [box('e', base), box('sub', lower), box('sup', upper)])


def frac(a, b):
    return box('f', [box('num', a), box('den', b)])


def summation(lower, upper, body):
    prop = el('m:naryPr'); prop.append(el('m:chr', '∑'))
    prop.append(el('m:limLoc', 'undOvr'))
    prop.append(el('m:grow', '1'))
    if not upper:
        prop.append(el('m:supHide', '1'))
    return box('nary', [prop, box('sub', lower), box('sup', upper), box('e', body)])


def equation_nodes(text):
    clean = re.sub(r'\s+', '', re.sub(r'\\tag\{\d+\}', '', text)).rstrip('.')
    expected = {
        1: r'\Delta_P=\frac{1}{K}\sum_{c=1}^{K}\frac{1}{|P_c|}\sum_{i\inP_c}(C_i^E-C_i^H)',
        2: r'\DeltaBA=G-L,\qquadT=G+L,\qquadQ=G+L+W',
        3: r'\DeltaT_S=T_S^E-T_S^H',
    }
    tag = re.search(r'\\tag\{(\d+)\}', text)
    number = int(tag.group(1)) if tag else None
    if number not in expected or clean != expected[number]:
        raise ValueError(f'Unrecognized display formula; extend reviewed OMML mapping first: {text}')
    r = mathrun
    if number == 1:
        difference = [r('('), sub([r('C')], [r('i')], [r('E')]), r('−', True),
                      sub([r('C')], [r('i')], [r('H')]), r(')')]
        inner = summation([r('i∈'), sub([r('P')], [r('c')])], [], difference)
        outer = summation([r('c=1')], [r('K')],
                          [frac([r('1', True)], [r('|'), sub([r('P')], [r('c')]), r('|')]), inner])
        return number, [sub([r('Δ')], [r('P')]), r('=', True), frac([r('1', True)], [r('K')]), outer]
    if number == 2:
        return number, [r('Δ'), r('BA', True), r('=G−L,     T=G+L,     Q=G+L+W')]
    return number, [r('Δ'), sub([r('T')], [r('S')]), r('=', True),
                    sub([r('T')], [r('S')], [r('E')]), r('−', True), sub([r('T')], [r('S')], [r('H')])]


def set_run_font(run, size=None, bold=None, italic=None):
    run.font.name = 'Times New Roman'
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'), 'Times New Roman'); fonts.set(qn('w:hAnsi'), 'Times New Roman')
    fonts.set(qn('w:eastAsia'), 'SimSun'); fonts.set(qn('w:cs'), 'Times New Roman')
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_inline(par, text, references, size=None, bold=False, italic=False, markup=True):
    def cite(m):
        keys = [v.strip() for v in m.group(1).split(',')]
        missing = set(keys) - references.keys()
        if missing: raise ValueError(f'Unknown citation keys {missing}')
        return '[' + ', '.join(str(references[k]['citation_number']) for k in keys) + ']'
    text = re.sub(r'\\cite\{([^{}]+)\}', cite, text)
    # Preserve each object label and its number on one rendered line.
    text = re.sub(r'\b(Figure|Fig\.|Table|Equation|Eq\.)\s+(\d+)', lambda m: m.group(1) + '\u00a0' + m.group(2), text)
    # URLs are literal identifiers, never inline mathematical notation.
    # Preserve underscores/carets before the generic sub/superscript parser.
    url_parts = re.split(r'(https?://[^\s]+)', text)
    if len(url_parts) > 1:
        for part in url_parts:
            if part.startswith(('https://', 'http://')):
                rr = par.add_run(part); set_run_font(rr, size, bold, italic)
            elif part:
                add_inline(par, part, references, size, bold, italic, markup)
        return
    if markup:
        tokens = re.split(r'(\*\*.*?\*\*|\*[^*\n]+\*|`[^`]+`)', text)
        for token in tokens:
            if not token: continue
            if token.startswith('**') and token.endswith('**'):
                add_inline(par, token[2:-2], references, size, True, italic, False)
            elif token.startswith('*') and token.endswith('*'):
                add_inline(par, token[1:-1], references, size, bold, True, False)
            elif token.startswith('`') and token.endswith('`'):
                rr = par.add_run(token[1:-1]); set_run_font(rr, size, bold, italic)
            else:
                add_inline(par, token, references, size, bold, italic, False)
        return
    if re.search(r'\\[A-Za-z]+|\$|\{\{|\}\}', text):
        raise ValueError(f'Unconverted LaTeX/placeholder in paragraph: {text}')
    pattern = re.compile(r'(?<![A-Za-z0-9])([A-Za-z]+|10)(?:_([A-Za-z0-9]+))?(?:\^(\{[^{}]+\}|[+−-]?\d+|[A-Za-z]))?')
    pos = 0
    for match in pattern.finditer(text):
        if match.group(2) is None and match.group(3) is None: continue
        plain = text[pos:match.start()]
        if plain:
            rr = par.add_run(plain); set_run_font(rr, size, bold, italic)
        rr = par.add_run(match.group(1)); set_run_font(rr, size, bold, italic)
        if match.group(2):
            rr = par.add_run(match.group(2)); set_run_font(rr, size, bold, italic); rr.font.subscript = True
        if match.group(3):
            rr = par.add_run(match.group(3).strip('{}').replace('-', '−'))
            set_run_font(rr, size, bold, italic); rr.font.superscript = True
        pos = match.end()
    if pos < len(text):
        rr = par.add_run(text[pos:]); set_run_font(rr, size, bold, italic)


def setup(doc):
    sec = doc.sections[0]
    sec.page_width = Mm(210); sec.page_height = Mm(297)
    sec.top_margin = Mm(18); sec.bottom_margin = Mm(18)
    sec.left_margin = Mm(18); sec.right_margin = Mm(18)
    sec.header_distance = Mm(8); sec.footer_distance = Mm(8)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'Footer']:
        s = doc.styles[name]; s.font.name = 'Times New Roman'; s.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(s.element.xpath('./w:pPr/w:pBdr')):
            border.getparent().remove(border)
        rp = s.element.get_or_add_rPr(); fonts = rp.get_or_add_rFonts()
        for theme_name in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme']:
            fonts.attrib.pop(qn('w:' + theme_name), None)
        for n in ['ascii', 'hAnsi', 'cs']: fonts.set(qn('w:' + n), 'Times New Roman')
        fonts.set(qn('w:eastAsia'), 'SimSun')
        s.font.size = Pt({'Title': 14, 'Heading 1': 11, 'Heading 2': 10, 'Heading 3': 10, 'Caption': 9}.get(name, 10))
        pf = s.paragraph_format; pf.space_after = Pt(5); pf.line_spacing = 1.1
        pf.widow_control = True
        if name.startswith('Heading'):
            s.font.bold = True; pf.keep_with_next = True; pf.space_before = Pt(10)
    doc.styles['Title'].font.bold = True
    doc.styles['Title'].paragraph_format.space_after = Pt(8)
    doc.styles['Caption'].font.italic = False
    footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = el('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    run = el('w:r'); t = el('w:t'); t.text = '1'; run.append(t); field.append(run); footer._p.append(field)
    settings = doc.settings.element
    update = el('w:updateFields', 'true'); settings.append(update)
    mathpr = el('m:mathPr'); mathpr.append(el('m:mathFont', 'Cambria Math')); settings.append(mathpr)


def add_authors(doc, refs):
    lines = [
        'Xin Li¹* · Bojian Guo¹ · Asel Kartanova¹',
        '¹ Higher School of Economics and Business, Department of Information Systems in Economics, I. Razzakov Kyrgyz State Technical University, 66 Ch. Aitmatov Ave., Bishkek 720044, Kyrgyz Republic',
        '* Corresponding author: Xin Li; lixin26@kstu.kg',
        'Xin Li: https://orcid.org/0009-0005-0670-5006; Bojian Guo: https://orcid.org/0009-0005-8353-2436; Asel Kartanova: https://orcid.org/0000-0003-1479-0747',
    ]
    for i, line in enumerate(lines):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        add_inline(p, line, refs, size=9, bold=(i == 0), markup=False)


def add_table(doc, rows, refs, table_number):
    rows = [r for r in rows if not all(re.fullmatch(r'\s*:?-+:?\s*', c) for c in r)]
    table = doc.add_table(rows=0, cols=len(rows[0])); table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [40, 70, 29, 35] if len(rows[0]) == 4 else [25, 41] + [27] * (len(rows[0])-2)
    factor = 174 / sum(widths); widths = [w * factor for w in widths]
    for col, width in zip(table.columns, widths): col.width = Mm(width)
    borders = el('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        edge = el('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
    table._tbl.tblPr.append(borders)
    # All current four manuscript tables are short enough to fit a full A4 page.
    # Keep the rows together so a single-page reader cannot miss continuation rows.
    keep_whole = table_number in {1, 2, 3, 4}
    for ri, values in enumerate(rows):
        row = table.add_row(); row._tr.get_or_add_trPr().append(el('w:cantSplit'))
        if ri == 0: row._tr.get_or_add_trPr().append(el('w:tblHeader', 'true'))
        for ci, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Mm(widths[ci]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr(); margin = el('w:tcMar')
            for side, amount in [('top', 65), ('bottom', 65), ('left', 75), ('right', 75)]:
                z = el('w:' + side); z.set(qn('w:w'), str(amount)); z.set(qn('w:type'), 'dxa'); margin.append(z)
            pr.append(margin)
            if ri == 0:
                shade = el('w:shd'); shade.set(qn('w:fill'), 'EEEEEE'); pr.append(shade)
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = (ri < len(rows)-1) if keep_whole else (ri == 0)
            p.paragraph_format.keep_together = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci < 2 else WD_ALIGN_PARAGRAPH.CENTER
            add_inline(p, value, refs, size=9, bold=(ri == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_formula(doc, expression):
    n, nodes = equation_nodes(expression)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.keep_together = True
    omp = el('m:oMathPara'); om = box('oMath', nodes + [mathrun('      (' + str(n) + ')', True)])
    omp.append(om); p._p.append(omp)


def add_figure(doc, number, captions, figures, refs):
    item = figures[number]
    path = Path(item['path'] if isinstance(item, dict) else item)
    width = item.get('width_mm', 174) if isinstance(item, dict) else 174
    # Keep image and separately editable caption text in one protected paragraph.
    # LibreOffice may ignore keepNext between an image-only and caption paragraph.
    p = doc.add_paragraph(style='Caption'); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.keep_with_next = False; p.paragraph_format.keep_together = True
    p.add_run().add_picture(str(path), width=Mm(width))
    p.add_run().add_break()
    add_inline(p, captions[number], refs, size=9)


def build(source, inventory, figures, output, language, expected_tables):
    text = source.read_text(encoding='utf-8').replace('\r\n', '\n')
    refs = {r['key']: r for r in inventory}
    parts = re.split(r'^## (?:Figure captions|图注)\s*$', text, maxsplit=1, flags=re.M)
    if len(parts) != 2: raise ValueError('Missing separate figure caption section')
    body, tail = parts
    captions = {m.group(1): m.group(0).strip() for m in re.finditer(r'^\*\*(?:Fig\.\s*|图\s*)([123])\*\*[^\n]+', tail, re.M)}
    if set(captions) != set(figures): raise ValueError('Figure caption/map mismatch')
    doc = Document(); setup(doc)
    doc.core_properties.title = body.splitlines()[0].lstrip('# ')
    doc.core_properties.author = 'Xin Li; Bojian Guo; Asel Kartanova'
    doc.core_properties.comments = 'Editable v1.9 manuscript from frozen scientific source'
    lines = body.splitlines(); i = 0; displayed = set(); formulas = 0; tables = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line: i += 1; continue
        if line == r'\[':
            content = []; i += 1
            while i < len(lines) and lines[i].strip() != r'\]': content.append(lines[i]); i += 1
            if i >= len(lines): raise ValueError('Unclosed display equation')
            add_formula(doc, '\n'.join(content)); formulas += 1; i += 1; continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append([x.strip() for x in lines[i].strip().strip('|').split('|')]); i += 1
            add_table(doc, rows, refs, tables + 1); tables += 1; continue
        heading = re.match(r'^(#{1,4})\s+(.*)', line)
        if heading:
            level, content = len(heading.group(1)), heading.group(2)
            p = doc.add_paragraph(style='Title' if level == 1 else 'Heading ' + str(level-1))
            add_inline(p, content, refs, bold=True)
            if level == 1: add_authors(doc, refs)
            i += 1; continue
        is_table_caption = bool(re.match(r'^\*\*(?:Table\s*|表\s*)\d+\*\*', line))
        p = doc.add_paragraph(style='Caption' if is_table_caption else None)
        if is_table_caption:
            p.paragraph_format.keep_with_next = True; p.paragraph_format.keep_together = True
        add_inline(p, line, refs, size=9 if is_table_caption else None)
        if not is_table_caption:
            for number in figures:
                if number not in displayed and re.search(r'(?:Figure\s*|Fig\.\s*|图\s*)' + number + r'(?!\d)', line):
                    add_figure(doc, number, captions, figures, refs); displayed.add(number)
        i += 1
    if displayed != set(figures): raise ValueError(f'Figures without body citation: {set(figures)-displayed}')
    resource = re.search(r'^\*\*(?:Online Resource 1|在线资源1)\*\*[^\n]+', tail, re.M)
    if resource:
        p = doc.add_paragraph(); add_inline(p, resource.group(0), refs)
    doc.add_paragraph('参考文献' if language == 'zh' else 'References', style='Heading 1')
    for ref in sorted(inventory, key=lambda r: r['citation_number']):
        p = doc.add_paragraph(); p.paragraph_format.left_indent = Mm(6); p.paragraph_format.first_line_indent = Mm(-6)
        p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_together = True
        add_inline(p, str(ref['citation_number']) + '. ' + ref['formatted'], refs, size=9, markup=False)
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)
    with ZipFile(output) as z:
        tree = etree.fromstring(z.read('word/document.xml'))
        omml = tree.xpath('count(//m:oMathPara)', namespaces=NS)
        formula_count = tree.xpath('count(//m:oMath)', namespaces=NS)
        table_count = tree.xpath('count(//w:tbl)', namespaces=NS)
        repeat_headers = tree.xpath('count(//w:tblHeader)', namespaces=NS)
        row_count = tree.xpath('count(//w:tr)', namespaces=NS)
        no_split = tree.xpath('count(//w:cantSplit)', namespaces=NS)
        native_text = ''.join(tree.xpath('//w:t/text()', namespaces=NS))
        media = [n for n in z.namelist() if n.startswith('word/media/')]
        if omml != 3 or formula_count != 3: raise ValueError('Need exactly three native OMML equations')
        if table_count != expected_tables or table_count != tables or repeat_headers != tables or no_split != row_count:
            raise ValueError('Native table structure check failed')
        if len(media) != 3: raise ValueError('Expected only three figure images, no equation images')
        if re.search(r'\\(?:cite|frac|Delta|tag|sum)|\*\*|`|\$|\{\{', native_text): raise ValueError('Raw markup leaked')
        expected_urls = re.findall(r'https?://[^\s]+', text + '\n' + '\n'.join(r['formatted'] for r in inventory))
        missing_urls = [url for url in expected_urls if url not in native_text]
        if missing_urls: raise ValueError(f'URL text was altered during formatting: {missing_urls}')
    return {'language': language, 'source': str(source), 'source_sha256': sha(source), 'docx': str(output),
            'docx_sha256': sha(output), 'native_omml_equations': int(omml), 'native_tables': int(table_count),
            'repeated_table_headers': int(repeat_headers), 'table_rows_protected_from_split': int(no_split),
            'figure_images': len(media), 'figure_numbers_inserted': sorted(displayed),
            'reference_count': len(inventory), 'raw_markup_check': 'PASS',
            'literal_source_and_reference_url_check': 'PASS', 'version': '1.9',
            'whole_table_keep_policy': [1, 2, 3, 4],
            'visual_review': 'NOT_YET_RENDERED'}


def build_cover_letter(source, output):
    """Format approved prose only; do not invent submission declarations."""
    content = source.read_text(encoding='utf-8').replace('\r\n', '\n').strip()
    doc = Document(); setup(doc)
    sec = doc.sections[0]
    sec.top_margin = Mm(23); sec.bottom_margin = Mm(23)
    sec.left_margin = Mm(23); sec.right_margin = Mm(23)
    for child in list(sec.footer._element): sec.footer._element.remove(child)
    doc.styles['Normal'].font.size = Pt(11)
    doc.styles['Normal'].paragraph_format.space_after = Pt(8)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.1
    doc.styles['Title'].font.size = Pt(13)
    doc.core_properties.author = 'Xin Li'
    doc.core_properties.title = 'Submission to Pattern Analysis and Applications'
    for block in re.split(r'\n\s*\n', content):
        block = block.strip()
        if not block: continue
        if block.startswith('# '):
            p = doc.add_paragraph(style='Title'); add_inline(p, block[2:], {}, bold=True)
        else:
            p = doc.add_paragraph(); add_inline(p, block.replace('\n', '\n'), {})
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)
    with ZipFile(output) as z:
        root = etree.fromstring(z.read('word/document.xml'))
        plain = ''.join(root.xpath('//w:t/text()', namespaces=NS))
        if re.search(r'\\(?:cite|frac|Delta|tag|sum)|\*\*|`|\$|\{\{', plain):
            raise ValueError('Raw markup leaked into cover letter')
    return {'language': 'cover_en', 'source': str(source), 'source_sha256': sha(source),
            'docx': str(output), 'docx_sha256': sha(output), 'version': '1.9',
            'visual_review': 'NOT_YET_RENDERED'}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--source-dir', type=Path, default=D19 / 'manuscript_revision')
    ap.add_argument('--output-dir', type=Path, default=D19 / 'manuscript_revision/native')
    ap.add_argument('--figure-map', type=Path, default=D19 / 'manuscript_revision/figure_map.json')
    ap.add_argument('--cover-letter', type=Path, required=True)
    ap.add_argument('--expected-tables', type=int, default=4)
    ap.add_argument('--final-source', action='store_true', help='Required assertion that the scientific sources have been frozen, not a readiness score')
    ap.add_argument('--render', action='store_true')
    ap.add_argument('--libreoffice', type=Path, help='Explicit absolute executable; no implicit installed-office fallback')
    ap.add_argument('--renderer-authorization-record', type=Path, help='Required if using a user-authorized non-bundled Windows executable')
    args = ap.parse_args()
    if not args.final_source:
        ap.error('Wait for source freeze, then provide --final-source. Old-source proofs are not supported.')
    inventory = json.loads((args.source_dir / 'reference_inventory.json').read_text(encoding='utf-8'))
    figures = json.loads(args.figure_map.read_text(encoding='utf-8'))
    reports = []
    for language, source_name in [('en', 'full_text.md'), ('zh', 'full_text_zh.md')]:
        stem = ('Pattern_Analysis_and_Applications_Research_Manuscript_v1.9' if language == 'en'
                else 'Pattern_Analysis_and_Applications_中文研究稿_v1.9')
        output = args.output_dir / (stem + '.docx')
        report = build(args.source_dir/source_name, inventory, figures, output, language, args.expected_tables)
        reports.append(report)
        print(json.dumps(report, ensure_ascii=True), flush=True)
    reports.append(build_cover_letter(args.cover_letter, args.output_dir / 'Cover_Letter.docx'))
    status = {'utc': datetime.now(timezone.utc).isoformat(), 'status': 'NATIVE_STRUCTURE_PASS_PENDING_VISUAL_QA',
              'version': '1.9', 'submission_readiness': 'NOT_CERTIFIED_BY_FORMAT_BUILDER',
              'builder_sha256': sha(__file__), 'sources_unchanged': True, 'reports': reports}
    (args.output_dir/'NATIVE_DOCUMENT_BUILD.json').write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding='utf-8')
    if args.render:
        if args.libreoffice is None or not args.libreoffice.is_absolute() or not args.libreoffice.is_file():
            raise RuntimeError('An explicit available absolute LibreOffice path is required; no fallback is permitted.')
        resolved_lo = args.libreoffice.resolve()
        is_bundled = resolved_lo.is_relative_to(RUNTIME.resolve())
        if not is_bundled:
            if not args.renderer_authorization_record or not args.renderer_authorization_record.is_file():
                raise RuntimeError('Non-bundled LibreOffice requires a recorded explicit user override of the skill prohibition.')
            status['renderer_authorization_record'] = str(args.renderer_authorization_record)
            status['renderer_authorization_sha256'] = sha(args.renderer_authorization_record)
        status['libreoffice_executable'] = str(resolved_lo)
        status['libreoffice_executable_sha256'] = sha(resolved_lo)
        env = os.environ.copy()
        env['PYTHONUTF8'] = '1'; env['PYTHONIOENCODING'] = 'utf-8'
        env['PYTHONDONTWRITEBYTECODE'] = '1'
        env['PATH'] = os.pathsep.join([str(resolved_lo.parent),
            str(RUNTIME/'dependencies/native/poppler/Library/bin'), env.get('PATH', '')])
        temp = args.output_dir/'renderer_temp'; temp.mkdir(parents=True, exist_ok=True)
        for k in ['TEMP', 'TMP', 'TMPDIR']: env[k] = str(temp)
        for report in reports:
            qa_dir = args.output_dir/'qa'/report['language']/report['docx_sha256'][:12]
            cmd = [sys.executable, str(SKILL/'render_docx.py'), report['docx'], '--output_dir',
                   str(qa_dir), '--emit_pdf']
            subprocess.run(cmd, env=env, check=True)
            from pypdf import PdfReader
            pdf = qa_dir/(Path(report['docx']).stem + '.pdf')
            count = len(PdfReader(pdf).pages)
            if len(list(qa_dir.glob('page-*.png'))) != count:
                raise RuntimeError('Renderer page-image count differs from current PDF')
            report.update(qa_dir=str(qa_dir), rendered_pdf=str(pdf), page_count=count,
                          visual_review='PENDING_FINAL_ALL_PAGE_INSPECTION')
        status['status'] = 'NATIVE_STRUCTURE_AND_RENDER_PASS_PENDING_VISUAL_QA'
        (args.output_dir/'NATIVE_DOCUMENT_BUILD.json').write_text(json.dumps(status, ensure_ascii=False, indent=2), encoding='utf-8')


if __name__ == '__main__': main()
